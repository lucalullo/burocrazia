from io import BytesIO

import fitz
import pytest
from docx import Document as DocxDocument

from src import ocr
from src.extract import (
    NO_TEXT_PDF_MESSAGE,
    NoTextFoundError,
    extract_docx_text,
    extract_pdf_text,
    extract_text,
    extract_txt_text,
)


def test_extract_txt() -> None:
    content = "Pratica amministrativa già registrata.".encode("utf-8")

    assert extract_txt_text(content) == "Pratica amministrativa già registrata."
    assert extract_text(content, "pratica.TXT").startswith("Pratica")


def test_extract_docx() -> None:
    document = DocxDocument()
    document.add_paragraph("Prima riga")
    document.add_paragraph("Seconda riga")
    buffer = BytesIO()
    document.save(buffer)

    assert extract_docx_text(buffer.getvalue()) == "Prima riga\nSeconda riga"


def test_extract_textual_pdf() -> None:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Documento amministrativo testuale")
    content = pdf.tobytes()
    pdf.close()

    assert "Documento amministrativo testuale" in extract_pdf_text(content)


def test_pdf_without_text_explains_ocr_limit(monkeypatch) -> None:
    pdf = fitz.open()
    pdf.new_page()
    content = pdf.tobytes()
    pdf.close()

    def unavailable(*_args, **_kwargs):
        raise ocr.OCRUnavailableError(ocr.OCR_UNAVAILABLE_MESSAGE)

    monkeypatch.setattr("src.extract.local_ocr.ocr_image", unavailable)

    with pytest.raises(NoTextFoundError, match="OCR") as exc_info:
        extract_pdf_text(content)

    assert str(exc_info.value) == NO_TEXT_PDF_MESSAGE
