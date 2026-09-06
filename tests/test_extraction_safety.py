from io import BytesIO
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
import zipfile

import fitz
from PIL import Image
from docx import Document
import pytest

from src import extract, ocr

TEXT = "Documento amministrativo sintetico per test locale"


def pdf_bytes(texts: list[str], *, password: str | None = None) -> bytes:
    with fitz.open() as pdf:
        for text in texts:
            page = pdf.new_page(width=300, height=200)
            if text:
                page.insert_text((20, 30), text)
        kwargs = {"encryption": fitz.PDF_ENCRYPT_AES_256, "owner_pw": password, "user_pw": password} if password else {}
        return pdf.tobytes(**kwargs)


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Prima")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Codice"
    table.cell(0, 1).text = "SR211"
    document.add_paragraph("Dopo")
    content = BytesIO()
    document.save(content)
    return content.getvalue()


def rewrite_zip(content: bytes, replacements: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(buffer, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info.filename, replacements.pop(info.filename, source.read(info)))
        for name, payload in replacements.items():
            # Assign after construction: ZipInfo otherwise sanitizes separators on Windows.
            info = ZipInfo()
            info.filename = info.orig_filename = name
            target.writestr(info, payload, compress_type=ZIP_DEFLATED)
    return buffer.getvalue()


def test_upload_stream_read_is_bounded(monkeypatch):
    monkeypatch.setattr(extract, "MAX_UPLOAD_BYTES", 32)
    stream = BytesIO(b"a" * 100)
    with pytest.raises(extract.ExtractionError, match="limite"):
        extract.extract_document(stream, "test.txt")
    assert stream.tell() == 33


@pytest.mark.parametrize("filename,content", [("test.pdf", b"plaintext"), ("test.docx", b"plaintext"), ("test.png", b"%PDF-1.7"), ("test.txt", b"%PDF-1.7\n")])
def test_camouflaged_formats_are_rejected(filename, content):
    with pytest.raises(extract.ExtractionError):
        extract.extract_document(content, filename)


def test_mime_mismatch_is_rejected():
    with pytest.raises(extract.ExtractionError, match="MIME"):
        extract.extract_document(TEXT.encode(), "test.txt", "application/pdf")
    assert extract.extract_document(TEXT.encode(), "test.txt", "text/plain; charset=utf-8").text == TEXT


def test_pdf_password_has_safe_error():
    with pytest.raises(extract.ExtractionError, match="password"):
        extract.extract_document(pdf_bytes([TEXT], password="synthetic-password"), "test.pdf")


def test_pdf_page_limit_is_enforced_before_ocr(monkeypatch):
    monkeypatch.setattr(extract, "MAX_PDF_PAGES", 2)
    monkeypatch.setattr(ocr, "ocr_image", lambda *_args, **_kwargs: pytest.fail("OCR oltre limite"))
    with pytest.raises(extract.ExtractionError, match="2 pagine"):
        extract.extract_document(pdf_bytes(["", "", ""]), "test.pdf")


def test_mixed_pdf_reads_scans_in_page_order(monkeypatch):
    monkeypatch.setattr(ocr, "ocr_image", lambda *_args, **_kwargs: ocr.OCRResult("Testo della scansione", "ita"))
    result = extract.extract_document(pdf_bytes(["Prima pagina digitale", "", "Terza pagina digitale"]), "test.pdf")
    assert result.text.splitlines() == ["Prima pagina digitale", "Testo della scansione", "Terza pagina digitale"]
    assert result.method == "pdf_mixed" and result.used_ocr
    assert (result.pages_processed, result.total_pages, result.ocr_pages) == (3, 3, 1)


def test_mixed_pdf_declares_pages_outside_ocr_coverage(monkeypatch):
    monkeypatch.setattr(ocr, "ocr_image", lambda *_args, **_kwargs: ocr.OCRResult(TEXT, "ita"))
    result = extract.extract_document(pdf_bytes([TEXT] * 5 + ["", TEXT]), "test.pdf")
    assert result.total_pages == 7 and result.pages_processed == 6
    assert "pagine: 6" in result.warning
    assert extract.PDF_OCR_LIMIT_WARNING in result.warning


def test_mixed_pdf_without_ocr_keeps_text_and_declares_gap(monkeypatch):
    def unavailable(*_args, **_kwargs):
        raise ocr.OCRUnavailableError(ocr.OCR_UNAVAILABLE_MESSAGE)
    monkeypatch.setattr(ocr, "ocr_image", unavailable)
    result = extract.extract_document(pdf_bytes([TEXT, ""]), "test.pdf")
    assert result.text == TEXT and result.pages_processed == 1
    assert "pagine: 2" in result.warning and "Tesseract" in result.warning


def test_pdf_raster_allocation_is_bounded(monkeypatch):
    monkeypatch.setattr(ocr, "MAX_OCR_IMAGE_PIXELS", 100_000)
    with fitz.open() as pdf:
        page = pdf.new_page(width=20_000, height=20_000)
        with extract._pdf_page_image(page) as image:
            assert image.width * image.height <= 100_000


@pytest.mark.parametrize("extension", ["txt", "pdf", "docx"])
def test_extracted_text_quota(extension, monkeypatch):
    monkeypatch.setattr(extract, "MAX_EXTRACTED_TEXT_CHARS", 5)
    content = {"txt": TEXT.encode(), "pdf": pdf_bytes([TEXT]), "docx": docx_bytes()}[extension]
    with pytest.raises(extract.ExtractionError, match="caratteri"):
        extract.extract_document(content, "test." + extension)


def test_docx_extracts_table_cells_in_body_order():
    assert extract.extract_document(docx_bytes(), "test.docx").text == "Prima\nCodice\nSR211\nDopo"


@pytest.mark.parametrize("constant,value", [("MAX_DOCX_MEMBERS", 2), ("MAX_DOCX_UNCOMPRESSED_BYTES", 100), ("MAX_DOCX_MEMBER_BYTES", 100), ("MAX_DOCX_XML_ELEMENTS", 2), ("MAX_DOCX_COMPRESSION_RATIO", 1)])
def test_docx_archive_quotas(constant, value, monkeypatch):
    monkeypatch.setattr(extract, constant, value)
    with pytest.raises(extract.ExtractionError):
        extract.extract_document(docx_bytes(), "test.docx")


@pytest.mark.parametrize("zip_separator", ["/", "\\"], ids=["posix_zip_separator", "windows_zip_separator_emulated"])
@pytest.mark.parametrize("name", ["../escape", "/absolute", "C:/absolute", "word\\escape"])
def test_docx_archive_rejects_unsafe_member_paths(name, zip_separator, monkeypatch):
    base = docx_bytes()
    # Exercise CPython's ZIP normalization independently of the host OS.
    with monkeypatch.context() as separators:
        separators.setattr(zipfile.os, "sep", zip_separator)
        content = rewrite_zip(base, {name: b"synthetic"})
        assert content.count(name.encode("ascii")) == 2  # Local and central headers retain the unsafe name.
        with ZipFile(BytesIO(content)) as archive:
            member = archive.infolist()[-1]
            assert member.orig_filename == name and archive.read(member) == b"synthetic"
        with pytest.raises(extract.ExtractionError, match="percorsi"):
            extract.extract_document(content, "test.docx")


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_docx_xml_entities_are_rejected(encoding):
    payload = '<?xml version="1.0"?><!DOCTYPE doc [<!ENTITY secret "synthetic">]><doc>&secret;</doc>'.encode(encoding)
    content = rewrite_zip(docx_bytes(), {"word/document.xml": payload})
    with pytest.raises(extract.ExtractionError, match="XML non consentite"):
        extract.extract_document(content, "test.docx")


def test_corrupt_docx_xml_is_reported_safely():
    content = rewrite_zip(docx_bytes(), {"word/document.xml": b"<broken>"})
    with pytest.raises(extract.ExtractionError, match="DOCX non è valido"):
        extract.extract_document(content, "test.docx")


def test_docx_duplicate_members_are_rejected():
    buffer = BytesIO(docx_bytes())
    with ZipFile(buffer, "a") as archive, pytest.warns(UserWarning):
        archive.writestr("word/document.xml", b"<document/>")
    with pytest.raises(extract.ExtractionError, match="duplicati"):
        extract.extract_document(buffer.getvalue(), "test.docx")


def test_image_pixel_limit_precedes_preprocessing(monkeypatch):
    monkeypatch.setattr(ocr, "MAX_INPUT_IMAGE_PIXELS", 10)
    monkeypatch.setattr(ocr.ImageOps, "exif_transpose", lambda *_args: pytest.fail("Decodifica oltre limite"))
    with Image.new("RGB", (10, 10)) as image, pytest.raises(ocr.ImageTooLargeError):
        ocr.prepare_image(image)


def test_ocr_timeouts_are_distinct_from_missing_component(monkeypatch):
    class Backend:
        @staticmethod
        def get_languages(**kwargs):
            assert kwargs["timeout"] == ocr.OCR_PROBE_TIMEOUT_SECONDS
            return ["ita"]
        @staticmethod
        def image_to_string(image, **kwargs):
            assert 0 < kwargs["timeout"] <= ocr.OCR_TIMEOUT_SECONDS
            raise ocr.OCRTimeoutError(ocr.OCR_TIMEOUT_MESSAGE)
    monkeypatch.setattr(ocr, "_configured_backend", lambda: Backend())
    with Image.new("RGB", (20, 20)) as image, pytest.raises(ocr.OCRTimeoutError, match="Tempo massimo"):
        ocr.ocr_image(image)


@pytest.fixture
def native_python_process(tmp_path, monkeypatch):
    """Adapt only the executable argv; run and observe a real child process."""
    command = tmp_path / "fake tesseract.py"
    processes, completed = [], []
    real_run, real_popen = ocr.subprocess.run, ocr.subprocess.Popen

    def observe_process(*args, **kwargs):
        process = real_popen(*args, **kwargs)
        processes.append(process)
        return process

    def run_python(args, **kwargs):
        assert args[0] == str(command)
        result = real_run([sys.executable, *args], **kwargs)
        completed.append(result)
        return result

    monkeypatch.setattr(ocr.subprocess, "Popen", observe_process)
    monkeypatch.setattr(ocr.subprocess, "run", run_python)
    monkeypatch.setattr(ocr, "TemporaryDirectory", lambda **kwargs: TemporaryDirectory(dir=tmp_path, **kwargs))
    return command, processes, completed


def test_native_timeout_removes_private_temporary_directory(tmp_path, native_python_process):
    command, processes, completed = native_python_process
    command.write_text("import time\nprint('synthetic-ready', flush=True)\ntime.sleep(30)\n", encoding="utf-8")
    with Image.new("RGB", (20, 20)) as image, pytest.raises(ocr.OCRTimeoutError) as error:
        ocr._LocalTesseractBackend(str(command)).image_to_string(image, lang="ita", timeout=2)
    cause = error.value.__cause__
    assert isinstance(cause, ocr.subprocess.TimeoutExpired) and b"synthetic-ready" in cause.output
    assert len(processes) == 1 and processes[0].returncode not in (None, 0)  # Killed and reaped by subprocess.run.
    assert not completed
    assert list(tmp_path.iterdir()) == [command]


def test_native_error_removes_private_temporary_directory(tmp_path, native_python_process):
    command, processes, completed = native_python_process
    command.write_text("print('synthetic-exit-two', flush=True)\nraise SystemExit(2)\n", encoding="utf-8")
    with Image.new("RGB", (20, 20)) as image, pytest.raises(ocr.OCRError, match="non riesce a elaborare") as error:
        ocr._LocalTesseractBackend(str(command)).image_to_string(image, lang="ita", timeout=5)
    assert type(error.value) is ocr.OCRError  # An executable-startup failure must not satisfy this test.
    assert len(processes) == len(completed) == 1 and processes[0].returncode == completed[0].returncode == 2
    assert b"synthetic-exit-two" in completed[0].stdout
    assert list(tmp_path.iterdir()) == [command]


def test_confident_orientation_wins_over_longer_garbage(monkeypatch):
    class Backend:
        calls = 0
        last_confidence = None
        @staticmethod
        def get_languages(**kwargs):
            return ["ita"]
        @staticmethod
        def image_to_osd(*args, **kwargs):
            raise RuntimeError("No OSD in fixture")
        def image_to_string(self, image, **kwargs):
            self.calls += 1
            self.last_confidence = 96 if self.calls == 1 else 12
            return TEXT if self.calls == 1 else "molto rumore inventato con numerose parole di scarsa qualità tutte errate"
    backend = Backend()
    monkeypatch.setattr(ocr, "_configured_backend", lambda: backend)
    with Image.new("RGB", (30, 20)) as image:
        result = ocr.ocr_image(image, auto_orient=True)
    assert result.text == TEXT and result.rotation == 0


def test_native_interruption_removes_temporary_directory(tmp_path, monkeypatch):
    monkeypatch.setattr(ocr, "TemporaryDirectory", lambda **kwargs: TemporaryDirectory(dir=tmp_path, **kwargs))
    def interrupted(*args, **kwargs):
        raise KeyboardInterrupt
    monkeypatch.setattr(ocr.subprocess, "run", interrupted)
    with Image.new("RGB", (20, 20)) as image, pytest.raises(KeyboardInterrupt):
        ocr._LocalTesseractBackend("synthetic-command").image_to_string(image, lang="ita", timeout=1)
    assert not list(tmp_path.iterdir())


def test_native_success_reads_confidence_and_cleans_files(tmp_path, native_python_process):
    command, processes, completed = native_python_process
    command.write_text(f"import sys\nfrom pathlib import Path\np=Path(sys.argv[2])\np.with_suffix('.txt').write_text({TEXT!r}, encoding='utf-8')\np.with_suffix('.tsv').write_text('level\\tpage_num\\tblock_num\\tpar_num\\tline_num\\tword_num\\tleft\\ttop\\twidth\\theight\\tconf\\ttext\\n5\\t1\\t1\\t1\\t1\\t1\\t0\\t0\\t20\\t10\\t97.5\\tDocumento\\n', encoding='utf-8')\n", encoding="utf-8")
    backend = ocr._LocalTesseractBackend(str(command))
    with Image.new("RGB", (20, 20)) as image:
        assert backend.image_to_string(image, lang="ita", timeout=5) == TEXT
    assert len(processes) == len(completed) == 1 and processes[0].returncode == completed[0].returncode == 0
    assert backend.last_confidence == 97.5
    assert list(tmp_path.iterdir()) == [command]


def test_pdf_timeout_is_an_extraction_error(monkeypatch):
    def timeout(*args, **kwargs):
        raise ocr.OCRTimeoutError(ocr.OCR_TIMEOUT_MESSAGE)
    monkeypatch.setattr(ocr, "ocr_image", timeout)
    with pytest.raises(extract.ExtractionError, match="Tempo massimo OCR"):
        extract.extract_document(pdf_bytes([""]), "test.pdf")
