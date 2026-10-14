from dataclasses import replace
import hashlib,json
from io import BytesIO
import pytest
from src.models import AdministrativeDocument
from src.document_fingerprints import binding
from src.fingerprint_v3 import FingerprintIndex,PARAMETERS,VERSION,features,page_record,validate,visual_similarity
from src.fingerprint_pages import RequestPages,image_signature,classify,field_signature,prepare
from src.extract import ExtractionResult,extract_document
from benchmarks.reference_truth import load,confirmed

TEXT='INPS\nMODELLO TEST-81 RICHIESTA SINTETICA DI ESEMPIO\nVersione 1.0\n'+('\nDati personali fittizi per la richiesta di esempio senza dati reali. '*12)

def document(did=9001):
    return AdministrativeDocument(id=did,source_id=2,ente='INPS',codice='TEST-81',titolo_ufficiale='RICHIESTA SINTETICA DI ESEMPIO',categoria=None,descrizione_ufficiale=None,significato_semplice=None,versione='1.0',data_rilevazione='2026-10-12',stato_verifica='da_verificare')

def record(d, texts=(TEXT,),sha=None,status='active'):
    return dict(document_id=d.id,source_id=2,catalog_binding_sha256=binding(d),canonical_file_sha256=sha or hashlib.sha256(str(d.id).encode()).hexdigest(),mime_type='application/pdf',url='https://example.gov.it/form.pdf',final_url='https://example.gov.it/form.pdf',acquisition_date='2026-10-12T00:00:00+00:00',generated_at='2026-10-12T00:00:00+00:00',size_bytes=123,status=status,shared_kind='unambiguous',pages=[page_record(i,t,t,()) for i,t in enumerate(texts,1)],aggregate=[],field_set=[],ocr_metadata={})

def index(records):return FingerprintIndex(dict(schema_version=3,fingerprint_version=VERSION,parameters=PARAMETERS,records=records))

def test_truth_registry_classifies_all_contradictions_and_preserves_locked_bytes():
    r=load();assert r['unclassified_contradictions']==0 and len(r['groups'])==9
    assert all(c['status']=='official_truth_confirmed' and c['label_status']=='confirmed' for c in confirmed())
    assert len(confirmed())==2
    assert all(c['label_status']!='confirmed' for c in r['cases'] if c['status']=='ambiguous')

def test_truth_refuses_unsupported_confirmation(tmp_path):
    r=load();r['cases'][0]['label_status']='confirmed'
    p=tmp_path/'truth.json';p.write_text(json.dumps(r))
    with pytest.raises(ValueError):load(p)

@pytest.mark.parametrize('n',['3','4','5'])
def test_stable_character_hash_and_exact_containment(n):
    a=features(TEXT);b=features(TEXT);assert a==b
    ix=index([record(document())]);rows=ix.similarities([a],n=n)
    assert rows[0]['chars']==1 and rows[0]['ocr']==1

def test_single_page_does_not_average_against_unrelated_canonical_pages():
    ix=index([record(document(),('Altra pagina completamente diversa senza modello. '*30,TEXT))])
    row=ix.similarities([features(TEXT)])[0]
    assert row['canonical_page']==2 and row['score']==1

def test_greedy_alignment_uses_distinct_input_and_canonical_pages():
    other=TEXT.replace('Dati personali','Informazioni ulteriori').replace('richiesta','istanza')
    ix=index([record(document(),(TEXT,other))]);r=ix.similarities([features(TEXT),features(other)],aggregation='top2')[0]
    assert len(r['alignment'])==2 and len({p[1] for p in r['alignment']})==2
    assert r['top2_average']<=r['max_similarity']

def test_ocr_to_ocr_remains_available_without_digital_text():
    d=document();r=record(d);r['pages']=[page_record(1,'',TEXT,())]
    row=index([r]).similarities([features(TEXT)])[0]
    assert row['digital']==0 and row['ocr']==1

def test_page_counters_and_full_scan_agree_on_supported_identity():
    ix=index([record(document()),record(document(9002),('Completamente estraneo privo di ancore. '*30,))]);q=[features(TEXT)];c={}
    a=ix.similarities(q,counters=c);b=ix.similarities(q,reference=True)
    assert a[0]==b[0] and c['candidates_visual']==0 and c['candidates_text']<=len(ix.pages)

def test_representation_classification():
    ex=ExtractionResult('test','pdf_text')
    assert classify('a.pdf',ex)=='digital_pdf'
    assert classify('a.pdf',replace(ex,used_ocr=True))=='scanned_pdf'
    assert classify('a.png',ex)=='image' and classify('a.docx',ex)=='docx'
    assert classify('a.pdf',ex,has_fields=True,filled=True)=='filled_acroform'
    assert classify('a.txt',ex)=='unknown'

def test_request_features_are_computed_once_without_global_input_cache(monkeypatch):
    import src.fingerprint_pages as pages
    original=pages.features;calls=[]
    monkeypatch.setattr(pages,'features',lambda t:(calls.append(t),original(t))[1])
    a=RequestPages((TEXT,));assert a.queries() is a.queries();assert calls==[TEXT]
    b=RequestPages((TEXT,));b.queries();assert calls==[TEXT,TEXT]

def test_visual_tiles_rotation_resize_and_crop():
    from PIL import Image,ImageDraw
    im=Image.new('RGB',(400,600),'white');draw=ImageDraw.Draw(im)
    for y in range(40,550,35):draw.rectangle((40,y,300+(y%60),y+8),fill='black')
    a=image_signature(im);assert len(a)==24
    with im.rotate(90,expand=True) as rotated:assert visual_similarity(a,image_signature(rotated))>.95
    with im.resize((200,300)) as small:assert visual_similarity(a,image_signature(small))>.90
    with im.crop((8,12,392,588)) as crop:assert visual_similarity(a,image_signature(crop))>.75

def test_decision_v4_lazy_visual_and_sha_safety(monkeypatch):
    import src.fingerprint_decision_v4 as decision
    d=document();ix=index([record(d)]);monkeypatch.setattr(decision,'public_index',lambda:ix)
    request=RequestPages((TEXT,),'digital_pdf',raw=b'not_needed')
    monkeypatch.setattr(request,'visual',lambda *args:pytest.fail('Digital decisive match rendered input'))
    result=decision.match(TEXT,[d],pages=request)
    assert result and result.document.id==d.id and request.counters['candidates_visual']==0
    assert decision.match(TEXT.replace('INPS','INAIL'),[d],file_sha256=ix.records[0]['canonical_file_sha256']) is None

def test_shared_resource_never_proves_identity_from_sha_alone(monkeypatch):
    import src.fingerprint_decision_v4 as decision
    a=document();b=document(9002);sha='a'*64
    ix=index([record(a,sha=sha,status='shared'),record(b,sha=sha,status='shared')]);monkeypatch.setattr(decision,'public_index',lambda:ix)
    assert decision.match(TEXT,[a,b],file_sha256=sha) is None

def test_explicit_revision_conflict_vetoes_content_match(monkeypatch):
    import src.fingerprint_decision_v4 as decision
    d=document();ix=index([record(d)]);monkeypatch.setattr(decision,'public_index',lambda:ix)
    assert decision.match(TEXT.replace('Versione 1.0','Versione 9.0'),[d]) is None

def test_rejected_record_and_stale_binding_are_inert(monkeypatch):
    import src.fingerprint_decision_v4 as decision
    d=document();ix=index([record(d)]);monkeypatch.setattr(decision,'public_index',lambda:ix)
    assert decision.match(TEXT,[replace(d,stato_verifica='rifiutato')]) is None
    assert decision.match(TEXT,[replace(d,titolo_ufficiale='Altro titolo')]) is None

def test_schema_rejects_unbounded_page_and_visual_data():
    d=document();b=dict(schema_version=3,fingerprint_version=VERSION,parameters=PARAMETERS,records=[record(d)])
    b['records'][0]['pages'][0]['visual']=['a'*16]*25
    with pytest.raises(ValueError):validate(b)

def static_pdf():
    import fitz
    with fitz.open() as pdf:
        p=pdf.new_page();p.insert_text((40,40),'INPS MODELLO TEST-81 RICHIESTA SINTETICA Versione 1.0');return pdf.tobytes()

def test_filled_static_reopens_with_all_synthetic_values_and_preserved_heading():
    import fitz
    from benchmarks.round8_fill import fill_static
    raw=static_pdf();out,audit=fill_static(raw,2)
    assert audit['reopened'] and audit['synthetic_verified'] and audit['fields_verified']==7
    with fitz.open(stream=out,filetype='pdf') as pdf:
        t=pdf[0].get_text();assert 'TEST-81' in t and 'FITTIZIA' in t and 'ZZZZZZ00A00Z000X' in t

@pytest.mark.parametrize('variant',[0,1])
def test_acroform_text_checkbox_persistence(variant):
    pytest.importorskip('pypdf');import fitz
    from benchmarks.round8_fill import fill_acroform
    with fitz.open() as pdf:
        p=pdf.new_page();p.insert_text((40,40),'MODULO FITTIZIO')
        for name,typ,rect in [('nome',fitz.PDF_WIDGET_TYPE_TEXT,(50,100,300,122)),('scelta',fitz.PDF_WIDGET_TYPE_CHECKBOX,(50,150,65,165))]:
            w=fitz.Widget();w.field_name=name;w.field_type=typ;w.rect=fitz.Rect(rect);p.add_widget(w)
        raw=pdf.tobytes()
    out,audit=fill_acroform(raw,variant)
    assert audit['fields_verified']==2 and audit['reopened']

def test_extraction_preserves_request_local_page_boundaries():
    import fitz
    with fitz.open() as pdf:
        for n in [1,2]:pdf.new_page().insert_text((40,40),f'Pagina {n} testo digitale di esempio sufficientemente lungo e leggibile.')
        raw=pdf.tobytes()
    ex=extract_document(raw,'a.pdf');assert ex.page_numbers==(1,2) and len(ex.page_texts)==2
    req=prepare(raw,'a.pdf',ex);assert len(req.queries())==2 and req.counters=={}

def test_pdf_checkbox_export_names_decode_hex_unicode():
    from benchmarks.round8_fill import pdf_state
    assert pdf_state('/S#EC')==pdf_state('/S#C3#AC')==pdf_state('/Sì')

def test_synthetic_values_follow_visible_label_types():
    from benchmarks.round8_fill import synthetic
    assert synthetic('EMAIL*',1).endswith('@example.invalid')
    assert synthetic('NATO/A IL GG/MM/AAAA',1)=='02/02/2030'
    assert synthetic('TELEFONO',1)=='0000000000'

def test_evaluation_single_shot_guard_precedes_recognition(monkeypatch,tmp_path):
    import benchmarks.round8_evaluation as ev
    monkeypatch.setattr(ev,'ROOT',tmp_path);monkeypatch.setattr(ev,'OUT','');monkeypatch.setattr(ev,'checked',lambda:{})
    (tmp_path/'evaluation_started.json').write_text('{}')
    monkeypatch.setattr(ev,'evaluate_worker',lambda *a:pytest.fail('Recognition must not restart'))
    with pytest.raises(ValueError,match='Single-shot'):ev.evaluate()

def test_evaluation_plan_cannot_be_reselected(monkeypatch,tmp_path):
    import benchmarks.round8_evaluation as ev
    monkeypatch.setattr(ev,'ROOT',tmp_path);monkeypatch.setattr(ev,'OUT','');monkeypatch.setattr(ev,'checked',lambda:{})
    (tmp_path/'robustness_v4_evaluation_plan.json').write_text('{}')
    with pytest.raises(ValueError,match='immutable'):ev.new_plan()

def test_frozen_hash_mismatch_blocks_evaluation(monkeypatch):
    import benchmarks.round8_common as common
    monkeypatch.setattr(common,'read',lambda p:dict(runtime_sha256='old',code_sha256='old',fingerprint_sha256='old'))
    monkeypatch.setattr(common,'identity',lambda:dict(runtime_sha256='new',code_sha256='new',fingerprint_sha256='new'))
    with pytest.raises(AssertionError,match='Frozen'):common.checked()

def test_input_feature_cache_does_not_mutate_public_index():
    ix=index([record(document())]);before=repr(ix.pages)
    for text in [TEXT,TEXT+' PERSONA FITTIZIA ZZZZZZ00A00Z000X']:
        request=RequestPages((text,));ix.similarities(request.queries())
        assert request.queries() is request.queries()
    assert repr(ix.pages)==before

def test_missing_revision_does_not_select_between_identical_peer_templates(monkeypatch):
    import src.fingerprint_decision_v4 as decision
    a=document();b=replace(document(9002),versione='2.0')
    text=TEXT.replace('Versione 1.0','')
    ix=index([record(a,(text,)),record(b,(text,))]);monkeypatch.setattr(decision,'public_index',lambda:ix)
    assert decision.match(text,[a,b]) is None

def test_metadata_batch_scorer_equals_reference_for_varied_local_input():
    from src.recognize import recognize_document
    docs=[document(),replace(document(9002),codice='TEST-82',titolo_ufficiale='RICHIESTA ALTRO ESEMPIO')]
    for text in [TEXT,TEXT.lower(),TEXT.replace('INPS','INAIL'),TEXT.replace('Versione 1.0','Versione 9.0'),'Modulo sconosciuto con nessun segnale di identità']:
        assert recognize_document(text,docs)==recognize_document(text,docs,reference_full_scan=True)

def test_uncertain_scan_renders_only_lazy_shortlist(monkeypatch):
    import src.fingerprint_decision_v4 as decision
    d=document();r=record(d);r['pages'][0]['visual']=['0'*16]*24
    ix=index([r]);monkeypatch.setattr(decision,'public_index',lambda:ix)
    request=RequestPages((TEXT,),'image',raw=b'local');calls=[]
    monkeypatch.setattr(request,'visual',lambda page:(calls.append(page) or ('0'*16,)*24))
    actual=decision.match(TEXT,[d],pages=request)
    assert actual and calls==[0] and request.counters['candidates_visual']==1

def test_late_title_code_and_entity_do_not_supply_identity_anchors(monkeypatch):
    import src.fingerprint_decision_v4 as decision
    d=document();ix=index([record(d)]);monkeypatch.setattr(decision,'public_index',lambda:ix)
    assert decision.match(('contenuto generico non identificativo '*70)+TEXT,[d]) is None

def test_prospective_selection_excludes_id_sha_url_and_peer_code(monkeypatch,tmp_path):
    import benchmarks.round8_evaluation as ev
    docs=[document(i) for i in range(9001,9006)];docs[3]=replace(docs[3],versione='2.0');docs[-1]=replace(docs[-1],codice='OTHER-99',titolo_ufficiale='Modulo indipendente con contenuto diverso')
    records=[record(d) for d in docs];records[1]['canonical_file_sha256']=records[0]['canonical_file_sha256']
    for i,r in enumerate(records):r['url']=r['final_url']='https://example.gov.it/'+str(i)+'.pdf'
    records[2]['url']=records[0]['url']
    monkeypatch.setattr(ev,'ROOT',tmp_path);monkeypatch.setattr(ev,'OUT','');monkeypatch.setattr(ev,'checked',lambda:dict(commit='frozen'))
    monkeypatch.setattr(ev,'catalog',lambda:(docs,[]));monkeypatch.setattr(ev,'calibration_ids',lambda:[9001]);monkeypatch.setattr(ev,'read_bundle',lambda:dict(records=records));monkeypatch.setattr(ev,'active_records',lambda:records);monkeypatch.setattr(ev,'read',lambda path:dict(cases=[]))
    captured={};monkeypatch.setattr(ev,'write',lambda path,value:captured.update({path:value}))
    ev.new_plan();assert captured['robustness_v4_evaluation_plan.json']['evaluation_identities']==[9005]
    leakage=captured['evaluation_v4_leakage.json'];assert leakage['prohibited_leakage']==0
    assert {r['reason'] for r in leakage['excluded_resources']}=={'identity_overlap','resource_sha_overlap','url_overlap','near_title_code_template'}

def test_request_score_and_entity_memo_is_local_and_rows_do_not_leak_visual(monkeypatch):
    import src.fingerprint_decision_v4 as decision
    d=document();r=record(d);r['pages'][0]['visual']=['0'*16]*24;ix=index([r]);monkeypatch.setattr(decision,'public_index',lambda:ix)
    calls=[];original=ix.similarities
    def counted(*a,**k):calls.append(1);return original(*a,**k)
    monkeypatch.setattr(ix,'similarities',counted)
    request=RequestPages((TEXT,),'image',raw=b'local');monkeypatch.setattr(request,'visual',lambda page:('0'*16,)*24)
    assert decision.match(TEXT,[d],pages=request)
    assert decision.match(TEXT,[d],pages=request)
    assert len(calls)==1 and request.counters['candidates_visual']==1
    cached=[v for k,v in request._memo.items() if k[0]=='page_scores'][0][0]
    assert all(r['visual']==0 for r in cached)
    decision.match(TEXT,[d],pages=RequestPages((TEXT,)))
    assert len(calls)==2

def test_metadata_fallback_cannot_override_dominant_content_with_another_id(monkeypatch):
    from types import SimpleNamespace
    from src.models import RecognitionResult,FamilyRecognitionResult
    import src.recognize as engine
    import src.fingerprint_decision_v7 as decision
    import src.fingerprint_v3 as fingerprints
    a=document();b=replace(document(9002),codice='OTHER-99')
    monkeypatch.setattr(fingerprints,'public_index',lambda:SimpleNamespace(records=[1]))
    def rejected(*args,**kwargs):kwargs['diagnostics']['dominant_content_candidate']=a.id;return None
    monkeypatch.setattr(decision,'match',rejected)
    monkeypatch.setattr(engine,'recognize_document',lambda *a,**k:RecognitionResult(True,95,document=b))
    monkeypatch.setattr(engine,'recognize_document_family',lambda *a,**k:FamilyRecognitionResult(False,0))
    result=engine.recognize_catalog(TEXT,[a,b],[])
    assert result.match_type=='NO_MATCH' and result.document is None
    assert 'metadata_content_conflict' in result.conflicts

def test_docx_fill_verifies_reopened_placeholder_text_only():
    from docx import Document
    from benchmarks.round8_fill import generate
    source=Document();source.add_paragraph('TITOLO UFFICIALE INPS');source.add_paragraph('Nome: _____________');source.add_table(rows=1,cols=1).cell(0,0).text='Note: ..............'
    buf=BytesIO();source.save(buf);out,audit=generate(buf.getvalue(),'form.docx',1);reopened=Document(BytesIO(out))
    assert audit['reopened'] and audit['synthetic_verified'] and audit['fields_verified']==2
    assert reopened.paragraphs[0].text=='TITOLO UFFICIALE INPS' and 'FITTIZIO' in reopened.paragraphs[1].text

def test_checkbox_only_template_uses_actual_synthetic_static_content():
    pytest.importorskip('pypdf');import fitz
    from benchmarks.round8_fill import generate
    with fitz.open() as pdf:
        p=pdf.new_page();p.insert_text((40,40),'MODELLO SINTETICO')
        w=fitz.Widget();w.field_name='scelta';w.field_type=fitz.PDF_WIDGET_TYPE_CHECKBOX;w.rect=fitz.Rect(50,100,65,115);p.add_widget(w);raw=pdf.tobytes()
    out,audit=generate(raw,'form.pdf',0)
    assert audit['acroform_static']=='static' and audit['fields_verified']==7
    with fitz.open(stream=out,filetype='pdf') as pdf:assert 'PERSONA FITTIZIA' in pdf[0].get_text()
